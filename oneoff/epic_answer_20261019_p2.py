# -*- coding: utf-8 -*-
"""Epic 2026-10-19 第5/6/8/9/10/11行 回答回填（行号+关键字双校验）"""
from openpyxl import load_workbook

PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"

ANSWERS = [
    (5, "多策略交集", (
        "已完成。选股结果页筛选栏新增「多策略交集」下拉（无默认值=不过滤，备选 ≥2策略/≥3策略/≥4策略），"
        "按同股同日命中策略数（strategies 聚合列）过滤。自测：≥2策略→143 只、≥3策略→11 只，"
        "与数据层全量计数一致，断言全过。"
    )),
    (6, "去掉这个", (
        "已完成。「多策略交集」独立模块整体下线：导航菜单项、page-inter 页面、交集 JS 全部函数、"
        "QP 交集实例、排序初始化、指数下拉交集副本全部移除，残余引用扫描为零；"
        "交集能力由选股结果页「多策略交集」筛选下拉承接；后端 /api/pick/intersection 接口保留可用。"
        "自测：导航 8 项无交集页，页面无 JS 报错。"
    )),
    (8, "涨跌幅", (
        "已完成。板块行情表格默认排序由流通市值改为涨跌幅降序（boardSort.defaultKeys=['pct_chg']），"
        "仍可点表头三态切换。自测：49 个板块首行=当日涨幅最大板块，降序断言通过。"
    )),
    (9, "暗盘信号", (
        "已完成。暗盘信号页改 50%:50% 分栏：左表在原有 10 列基础上新增 命中策略/基准日/连涨天数/连涨涨幅 "
        "4 列（后端 /api/darkpool/signals 补 stock_daily_quote 的 cont_days/cont_pct + 同日选股命中策略聚合，"
        "与选股结果页同口径）；列宽随内容/列名自适应（table-layout:auto + nowrap），展示不全出横向滚动条；"
        "全 14 列支持三态自定义排序，默认涨跌幅降序；右侧行情K线（QP 面板）随左表点行联动。"
        "自测：分栏/14列/默认降序断言通过，点击 301313.SZ 右侧 K 线加载成功。"
    )),
    (10, "涨停池", (
        "已完成。涨停池页改 50:50：左侧 50% 靠左为现有内容（筛选+连板梯队/行业分布图+表格），"
        "表格列宽自适应+横向滚动；右侧 50% 新增行情K线（QP 面板），随左表点行联动（自测 003029.SZ 加载成功）。"
        "另修复小视口下左栏内容溢出：分栏面板体加纵向滚动兜底，梯队/行业图高度 200→170px。"
    )),
    (11, "市场面板", (
        "已完成三项：①人气榜与资金流对比模块等高（panel-body 430px），人气榜超出部分上下滚动；"
        "人气榜加 TOP 单选筛选（TOP10/TOP15/TOP20，默认 TOP10）；"
        "②资金流对比的净流入/净流出加 TOP 单选筛选（TOP5/TOP10/TOP15/TOP20，默认 TOP5）；"
        "③页面最底部新增通栏「行情K线」模块（QP 面板），点击人气榜行或资金流股票即联动展示。"
        "自测：默认 10 行/5 条，切 TOP20→20 行、资金流 TOP10→10 条，点 600487.SH 底部 K 线加载成功，"
        "两模块高度均 430px。"
    )),
]

wb = load_workbook(PATH)
ws = wb["Sheet1"]
for r, kw, ans in ANSWERS:
    detail = ws.cell(row=r, column=3).value or ""
    if kw not in detail:
        raise SystemExit(f"[中止] 第{r}行关键字 {kw!r} 校验失败: {detail[:60]!r}")
for r, kw, ans in ANSWERS:
    ws.cell(row=r, column=5).value = ans
    print(f"[回填] 第{r}行 {len(ans)} 字")
wb.save(PATH)
print("已保存")
