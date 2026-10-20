# -*- coding: utf-8 -*-
"""回填 epic_description.xlsx 回答列（2026-10-14 任务：选股策略来源调研）
按行号 + 关键字双重校验，与 update_epic_status.py 同一套防误伤约定。
"""
import sys
from openpyxl import load_workbook

PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"

ROW = 2          # 目标行
KEYWORD = "选股策略"   # 必须出现在该行"任务详情"中的关键字
ANSWER = (
    "策略全部为本系统自研，与 akshare 无关。\n"
    "① 来源：app/strategy/ 目录下 13 个策略类（MA_BULL 均线多头 / VOLUME_BREAKOUT 放量突破 / "
    "VALUE 低估值 / MACD_GOLDEN MACD金叉 / PULLBACK 回踩 / CONSECUTIVE_YANG 连续阳线 / "
    "WEEKLY_MAINLINE 每周主线 / ZT_POOL_FOLLOW 涨停跟随 / VALUE_PCTL 估值分位 / "
    "FUND_FLOW_DIVERGE 资金流背离 / MARGIN_TREND 两融趋势 / VALUE_QUALITY 质优低估 / "
    "FIN_QUALITY 业绩质量），业务思想主要来自 docs/03-一周策略.md 与 docs/04-一周策略-量化实现.md "
    "的量化实现。akshare 在本系统中仅作数据采集源（app/collector/akshare_collector.py，vendored 源码"
    "接入见 config.AKSHARE_PATH，采集涨停池/强势股池/估值历史/业绩报表等 stock_* 数据），"
    "它不提供任何选股策略。\n"
    "② 实现机制：BaseStrategy 抽象基类 + STRATEGY_REGISTRY 注册表模式（app/strategy/base.py）。"
    "每个策略一个 Python 类，声明 code/name/default_params 并实现 select()，模块内 register() 自注册；"
    "运行时 get_strategy(code) 按编码取用，前端经 /api/strategies 列出、"
    "/api/pick|/api/pick/asof|/api/pick/live 按 code 调用。\n"
    "③ 写死 vs 配置表：半写死。策略集合靠 pick_service.py 顶部写死的 import 列表触发注册"
    "（L15-27 逐个 import app.strategy.xxx）——新增策略必须① 新建 app/strategy/xxx.py ② 在 "
    "pick_service.py 加一行 import，无自动发现机制；数据库中不存在策略配置表（schema.sql 全部 33 张表"
    "中无策略定义表，stock_pick_batch/stock_pick_result 仅存执行批次与命中明细）。但调用侧是动态的："
    "API 可按任意已注册 code 调用，参数支持 req.params 运行时覆盖 default_params（merged_params 合并），"
    "执行参数以 JSON 留痕于 stock_pick_batch.params（是记录，不是配置来源）。"
)

wb = load_workbook(PATH)
ws = wb["Sheet1"]
detail = ws.cell(row=ROW, column=3).value or ""
if KEYWORD not in detail:
    print(f"[校验失败] 第{ROW}行任务详情未包含关键字 {KEYWORD!r} (实际: {detail[:60]!r})")
    sys.exit(1)

old = ws.cell(row=ROW, column=5).value
if old:
    print(f"[提示] 第{ROW}行回答列已有内容, 将覆盖 (旧值前60字: {str(old)[:60]!r})")
ws.cell(row=ROW, column=5).value = ANSWER

try:
    wb.save(PATH)
    print(f"[已写入] 第{ROW}行 回答列 ({len(ANSWER)} 字) -> {PATH}")
except PermissionError as e:
    print(f"[失败] 文件被占用(请关闭 Excel): {e}")
    sys.exit(3)
