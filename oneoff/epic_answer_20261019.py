# -*- coding: utf-8 -*-
"""Epic 2026-10-19 三行任务回答回填（行号+关键字双校验, 与 update_epic_status 同规约）"""
from openpyxl import load_workbook

PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"

ANSWERS = [
    (2, "选股策略", (
        "可以，已完成配置化改造（v1.6）。新建 stock_strategy_config 表"
        "（strategy_code 唯一 / enabled 启停 / params_json 参数覆盖 / remark），"
        "合并优先级：类内 default_params < 库表 params_json < 运行时传入；"
        "服务启动自动建表播种（seed=类默认值，不改变任何策略行为）。"
        "run_pick / pick_asof / pick_live / 回测全部统一读库表口径；"
        "enabled=0 时 run_pick_all（每日盘后调度）跳过该策略。"
        "新增接口：GET /api/strategy_config 查看全量（默认值/覆盖/生效值），"
        "POST /api/strategy_config/{code} 改参数或停用，即时生效、无需改代码重启。"
        "自测：MA_BULL ma_rising_days 3→5 后 asof 命中 221→199；"
        "运行时传参优先级高于库表（回到 221）；启停翻转正常；恢复默认后 221。"
        "代码：app/strategy/config_store.py + pick_service.py/backtest.py/api main.py 接线。"
    )),
    (3, "akshare", (
        "审计结论：不是全部，也无必要全部实现。vendored akshare 共 410 个 stock_* 函数，"
        "本系统按「数据域」覆盖而非按函数逐个接入——已落库 39 张 stock_* 表，"
        "直接引用 akshare 函数 8 个（涨停/强势/跌停/炸板/昨涨停/估值历史/业绩报表/资金流），"
        "其余同主题数据由新浪/腾讯/东财直连/Baostock 通道覆盖（往往更稳更快）。"
        "已覆盖域：日线/分钟行情、估值全历史+分位、资金流(新浪+THS双源)、两融、龙虎榜、大宗、"
        "业绩报表、业绩预告、商誉、质押、股东户数、解禁、停复牌、筹码、人气榜、板块/概念/申万行业、复权因子。"
        "本次补缺开发：董监高持股变动 stock_share_change（东财 RPT_EXECUTIVE_HOLD_DETAILS，"
        "全量 172,525 行已落库、覆盖 4547 只股票，增量幂等已验证，进 --all 每日链路，"
        "--sharechange [--full]）。"
        "仍未覆盖（诚实清单）：①三大报表明细(利润/现金/资产负债，单股接口全市场回填约1.5万请求，按需再补)；"
        "②十大股东/十大流通股东(单股接口)；③机构持仓明细；④北向持股明细(hsgt，2024-08 起交易所停止披露，不可采)；"
        "⑤港股/美股/基金/ESG/宏观类(与 A 股选股无关)。策略需要哪个再按域补哪个。"
    )),
    (4, "缠论", (
        "已实现并注册：CHAN_BUY2 缠论二买（第 14 个策略，app/strategy/chan_buy2.py）。"
        "简化缠论日线量化：窗口内底分型一买 L1 → 反弹≥5% → 回踩不破一买(L2≥L1)且幅度2%~15%、"
        "L2 为底分型 → 3 日内收阳确认；缩量回踩/DIF 底背驰(pandas 现算)/站上 MA5 加分；"
        "K 线包含关系不处理（明示简化项）。参数已入 stock_strategy_config 可在线调整。"
        "自测：2026-10-19 asof 命中 50 只(top_n 截断)，命中原因含 L1/反弹/回踩/缩量/背驰明细。"
        "回测诚实口径：10-12 基准日 50 命中、持有 5 日胜率 18%、平均 -3.63%"
        "（同期市场回调，二买在下跌中继失效属预期风险）；纪律化退出(-3%止损)平均 -2.45%。"
        "单基准日小样本，按项目约定不据此调参（防过拟合）。"
    )),
]

wb = load_workbook(PATH)
ws = wb["Sheet1"]
ok = True
for r, kw, ans in ANSWERS:
    detail = ws.cell(row=r, column=3).value or ""
    if kw not in detail:
        print(f"[中止] 第{r}行关键字 {kw!r} 校验失败: {detail[:60]!r}")
        ok = False
        break
if ok:
    for r, kw, ans in ANSWERS:
        ws.cell(row=r, column=5).value = ans
        print(f"[回填] 第{r}行 回答 {len(ans)} 字")
    try:
        wb.save(PATH)
        print("已保存")
    except PermissionError as e:
        print(f"[失败] 文件被占用(请关闭 Excel): {e}")
