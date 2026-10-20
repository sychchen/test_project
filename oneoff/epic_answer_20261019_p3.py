# -*- coding: utf-8 -*-
"""Epic 第10行回答补充：涨停池「首板」筛选子项（用户中途扩写的第①点）。"""
import io
import openpyxl

PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"

ANSWER_10 = (
    "已完成两项：①筛选条件新增「首板」选项（精确 cont_boards=1），选项排序为 "
    "全部涨停 / 首板 / ≥2连板 / ≥3连板 / ≥4连板 / ≥5连板；后端 /api/ztpool/pool 新增 "
    "exact_boards 参数（精确匹配优先于 min_boards 下限过滤，梯队/行业/明细三处 SQL 同口径）。"
    "自测：2026-10-19 首板 35 只全部 cont_boards=1（非首板 0 行），全部涨停 46 只不受影响；"
    "浏览器实测下拉选项顺序正确、首板筛选 35 行、梯队图仅首板一根柱。"
    "②涨停池页改 50:50：左侧 50% 靠左为现有内容（筛选+连板梯队/行业分布图+表格），"
    "表格列宽自适应+横向滚动，右侧挂行情K线（QP zt 实例），点击表格行联动查看个股K线。"
)

wb = openpyxl.load_workbook(PATH)
ws = wb.active
row = 10
detail = str(ws.cell(row=row, column=3).value or "")
assert "涨停池" in detail and "首" in detail, f"第{row}行关键字校验失败: {detail[:50]}"
ws.cell(row=row, column=5).value = ANSWER_10
wb.save(PATH)
print(f"OK 第{row}行回答已更新（含首板筛选说明）")
