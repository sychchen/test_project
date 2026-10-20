# -*- coding: utf-8 -*-
"""
read_epic.py — 固定脚本：直接从磁盘读取 epic_description.xlsx（无缓存）

用法：
    python read_epic.py [xlsx路径]

说明：
    - 每次运行都重新从磁盘加载工作簿（openpyxl 直接读文件，不经过 editor_sdk
      等任何内存缓存），保证读到的是最新保存的内容。
    - 输出：工作簿结构（sheet / 行列规模 / 表头）+ 全部数据行完整内容。
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(DEFAULT_PATH)
    if not path.exists():
        print(f"[ERROR] 文件不存在: {path}")
        sys.exit(1)

    # data_only=True 读公式计算后的值；每次调用都是全新加载，无缓存
    wb = load_workbook(path, data_only=True, read_only=True)

    print(f"文件: {path}")
    print(f"修改时间: {__import__('datetime').datetime.fromtimestamp(path.stat().st_mtime):%Y-%m-%d %H:%M:%S}")
    print(f"Sheet 数: {len(wb.sheetnames)} -> {wb.sheetnames}")
    print("=" * 70)

    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
        # 去掉全空行，定位真实数据区
        nonempty = [r for r in rows if any(v.strip() for v in r)]
        print(f"\n## Sheet: {ws.title}  (声明范围 {ws.max_row} 行 x {ws.max_column} 列, "
              f"实际非空 {len(nonempty)} 行)")
        if not nonempty:
            print("  (空表)")
            continue
        width = max(len(r) for r in nonempty)
        for i, r in enumerate(nonempty, 1):
            cells = [v for v in r] + [""] * (width - len(r))
            # 只展示到最后一个非空列
            last = max((j for j, v in enumerate(cells) if v.strip()), default=-1)
            cells = cells[: last + 1]
            tag = "表头" if i == 1 else f"第{i}行"
            print(f"  [{tag}] " + " | ".join(cells))

    wb.close()


if __name__ == "__main__":
    main()
