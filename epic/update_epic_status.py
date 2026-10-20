# -*- coding: utf-8 -*-
"""回填 epic_description.xlsx 状态列(按精确行号 + 全文本关键字双重校验, 防误伤新增行)"""
import sys
from openpyxl import load_workbook

PATH = r"D:\kiro-worspace\test_project\stock-system\epic\epic_description.xlsx"

# (行号, 状态, 必须出现在该行"任务详情"中的关键字)
# 用法: 每次处理前先用 read_epic.py 确认行号, 只保留"本次要写"的行(历史行会让校验输出变噪音)
ROWS = [
    (7,  "已完成", "板块行情"),        # 板块分时断更根因修复+历史回填 (2026-10-20)
]

wb = load_workbook(PATH)
ws = wb["Sheet1"]
print(f"维度: {ws.max_row} 行 x {ws.max_column} 列")

changes = []
for r, status, kw in ROWS:
    detail = ws.cell(row=r, column=3).value or ""
    old = ws.cell(row=r, column=1).value
    if kw not in detail:
        print(f"[跳过] 第{r}行 校验失败: 未包含关键字 {kw!r} (实际: {detail[:60]!r})")
        continue
    if old == status:
        print(f"[无需改动] 第{r}行 已经是 {status}")
        continue
    ws.cell(row=r, column=1).value = status
    changes.append((r, old, status, detail[:50]))

if not changes:
    print("无变更，退出")
    sys.exit(0)

for r, o, n, d in changes:
    print(f"[改动] 第{r}行 {o} -> {n}  | {d}")
try:
    wb.save(PATH)
    print(f"已保存 {len(changes)} 处变更 -> {PATH}")
except PermissionError as e:
    print(f"[失败] 文件被占用(请关闭 Excel): {e}")
    sys.exit(3)
